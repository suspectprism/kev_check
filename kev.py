# Check for updates to CISA KEV (Known Exploited Vulnerabilities) list
#
# v0.1 Download CISA KEV list from CISA website and extract relevant data
#      Save updated KEV data to local Excel spreadsheet
#
# v0.2 If KEV list is updated then notify to a private Discord channel
#
# v0.3 Support use of a generic (instead of hard-coded) vendor when checking for vendor vulns
#      Improved formatting of Discord notification message
#      List new vulnerabilities in Discord message
#
# v0.4 Compare new vulns against previous stored vulns, for reporting to Discord.
#      Allows for CISA changing the order of reported vulnerabilities in the KEV list.
#
# v0.5 Optionally download VulnCheck KEV list from VulnCheck API and save updated data to the spreadsheet.
#
# v0.5.1 Notify to another Discord channel if there are changes to the VulnCheck KEV list
#
# P Dowley   v0.5.1      29 Mar 2026

import requests
from discord_webhook import DiscordWebhook, DiscordEmbed
import sys
from pathlib import Path
import openpyxl
from openpyxl import styles
from openpyxl import utils
from datetime import date, datetime
import shutil
from colorama import Fore, Style
from config import load_config

def get_cisa_kev_data(kev_url):
    '''Fetch KEV data from the CISA website'''
    try:
        response = requests.get(kev_url)
        response.raise_for_status()  # Raise an error for bad status codes
        kev_data = response.json()  # Parse the JSON content into a Python dictionary
        print("CISA KEV data downloaded successfully.")
        return kev_data
    except requests.exceptions.RequestException as e:
        print(f"Error retrieving CISA KEV data from {kev_url}: {e}")
        sys.exit(1)  # Exit on error with a non-zero status

def get_vc_kev_data(vc_url, vc_token):
    '''Fetch all KEV data from the VulnCheck API using cursor-based pagination'''
    try:
        headers = {"Authorization": f"Bearer {vc_token}"}
        all_data = []
        meta = {}

        # First page
        params = {"start_cursor": "true", "limit": 500}
        response = requests.get(vc_url, headers=headers, params=params)
        response.raise_for_status()
        result = response.json()
        all_data.extend(result.get("data") or [])
        meta = result.get("_meta") or result.get("meta") or {}

        # Subsequent pages
        while meta.get("next_cursor"):
            params = {"cursor": meta["next_cursor"]}
            response = requests.get(vc_url, headers=headers, params=params)
            response.raise_for_status()
            result = response.json()
            all_data.extend(result.get("data") or [])
            meta = result.get("_meta") or result.get("meta") or {}

        print("VulnCheck KEV data downloaded successfully.")

        return {"_meta": meta, "data": all_data}
    except requests.exceptions.RequestException as e:
        print(f"Error retrieving VulnCheck KEV data from {vc_url}: {e}")
        sys.exit(1)

def write_vc_sheets(wb, vc_meta, vc_data, summary_ws, vendor_list, vendor_name):
    '''Add VC summary rows to the Summary sheet and add/replace VC_Vulns and VC_<vendor> sheets'''

    # Remove existing VC sheets if present (e.g. on a refresh run)
    vc_vend_sheet = "VC_" + vendor_name
    for sheet_name in ("VC_Vulns", vc_vend_sheet):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    # Write VC metadata into the Summary sheet at rows 7-9 (row 6 is a blank separator)
    vc_summary_rows = [
        (7, "VC Title",                    vc_meta.get("index", "")),
        (8, "VC Timestamp",                vc_meta.get("timestamp", "")),
        (9, "VC Count of Vulnerabilities", len(vc_data)),
    ]
    for row_num, label, value in vc_summary_rows:
        summary_ws.cell(row=row_num, column=1).value = label
        cell_b = summary_ws.cell(row=row_num, column=2)
        cell_b.value = value
        cell_b.font = styles.Font(bold=True)
    summary_ws['B9'].alignment = styles.Alignment(horizontal='left')

    # VC_Vulns sheet
    vc_vulns_ws = wb.create_sheet(title="VC_Vulns")

    headers = [
        "CVE ID", "Vendor", "Product", "Vulnerability Name", "CWEs",
        "Date Added (VC)", "CISA Date Added", "Due Date",
        "Short Description", "Required Action",
        "Ransomware Campaign", "Canary Exploitation",
        "XDB Exploit URLs", "Reported Exploitation URLs",
    ]
    vc_vulns_ws.append(headers)

    for cell in vc_vulns_ws[1]:
        cell.font = styles.Font(bold=True)

    col_widths = [15, 15, 20, 40, 15, 12, 12, 12, 40, 40, 20, 20, 50, 50]
    for i, width in enumerate(col_widths, start=1):
        vc_vulns_ws.column_dimensions[utils.get_column_letter(i)].width = width

    def _iso_date(value):
        '''Parse a timestamp string and return a yyyy-mm-dd string, or "" if empty'''
        if not value:
            return ""
        try:
            return datetime.fromisoformat(str(value).replace("Z", "+00:00")).strftime("%Y-%m-%d")
        except ValueError:
            return str(value)

    for item in vc_data:
        cve_ids = ", ".join(item.get("cve") or [])
        cwes    = ", ".join(item.get("cwes") or [])
        xdb_urls = ", ".join(
            e.get("xdb_url", "") for e in (item.get("vulncheck_xdb") or [])
        )
        exploitation_urls = ", ".join(
            e.get("url", "") for e in (item.get("vulncheck_reported_exploitation") or [])
        )

        row = [
            cve_ids,
            item.get("vendorProject", ""),
            item.get("product", ""),
            item.get("vulnerabilityName", ""),
            cwes,
            _iso_date(item.get("date_added")),
            _iso_date(item.get("cisa_date_added")),
            _iso_date(item.get("dueDate")),
            item.get("shortDescription", ""),
            item.get("required_action", ""),
            item.get("knownRansomwareCampaignUse", ""),
            item.get("reported_exploited_by_vulncheck_canaries", ""),
            xdb_urls,
            exploitation_urls,
        ]
        vc_vulns_ws.append(row)

    vc_vulns_ws.auto_filter.ref = vc_vulns_ws.dimensions

    # VC vendor sheet — filtered copy of VC_Vulns
    vc_vend_ws = wb.copy_worksheet(vc_vulns_ws)
    vc_vend_ws.title = vc_vend_sheet

    for row_ctr in range(vc_vend_ws.max_row, 1, -1):
        vendor_cell = vc_vend_ws.cell(row=row_ctr, column=2)  # Vendor is in the second column
        if vendor_cell.value not in vendor_list:
            vc_vend_ws.delete_rows(row_ctr)

    vc_vend_ws.auto_filter.ref = vc_vend_ws.dimensions

    # Write VC vendor count into Summary row 10
    vc_vend_count = vc_vend_ws.max_row - 1  # Exclude header row
    vc_vend_count_header = "VC Count of " + vendor_name + " vulns"
    summary_ws.cell(row=10, column=1).value = vc_vend_count_header
    cell_b10 = summary_ws.cell(row=10, column=2)
    cell_b10.value = vc_vend_count
    cell_b10.font = styles.Font(bold=True)
    cell_b10.alignment = styles.Alignment(horizontal='left')

def check_vc_updates(kev_in_path, vc_data, vendor_list, vendor_name, notify_discord=False, webhook_url=None):
    '''Check if VulnCheck KEV data has changed since the last saved file'''
    wb = openpyxl.load_workbook(kev_in_path)

    # VC sheets have never been saved
    if "VC_Vulns" not in wb.sheetnames:
        print("VulnCheck KEV data: not previously saved.")
        return True

    # VC rows may be absent if file was saved before VC support was added
    summary_ws = wb["Summary"]
    if summary_ws.cell(row=9, column=1).value != "VC Count of Vulnerabilities":
        print("VulnCheck KEV data: Summary rows not present in saved file.")
        return True

    saved_count = summary_ws.cell(row=9, column=2).value or 0
    saved_vend_count = summary_ws.cell(row=10, column=2).value or 0
    current_count = len(vc_data)
    current_vend_count = sum(1 for v in vc_data if v.get("vendorProject") in vendor_list)

    print(f" Saved VulnCheck KEV count: {saved_count}, {vendor_name}: {saved_vend_count}")

    if saved_count != current_count or saved_vend_count != current_vend_count:
        count_str = f"{Fore.YELLOW}{current_count}{Style.RESET_ALL}" if saved_count != current_count else str(current_count)
        vend_str = f"{Fore.RED}{current_vend_count}{Style.RESET_ALL}" if saved_vend_count != current_vend_count else str(current_vend_count)
        print(f"Latest VulnCheck KEV count: {count_str}, {vendor_name}: {vend_str}")

        if notify_discord and webhook_url:
            saved_cve_ids = set()
            for row in wb["VC_Vulns"].iter_rows(min_row=2, values_only=True):
                if row[0]:
                    saved_cve_ids.update(c.strip() for c in str(row[0]).split(","))
            notify_vc_to_discord(saved_count, saved_vend_count, vc_data,
                                  saved_cve_ids, vendor_name, current_vend_count, webhook_url)
        else:
            print("Notification to Discord is not required.")

        return True
    else:
        print(f"Latest VulnCheck KEV count: {current_count}, {vendor_name}: {current_vend_count} (no change)")
        return False

def notify_to_discord(saved_summary_dict, kev_data, vulns_list, vend_name, vend_count, webhook_url, saved_cve_ids):
    '''Send message to private Discord channel via webhook to notify of KEV updates'''

    # Convert string with ISO format date to a datetime
    saved_release_dt = datetime.fromisoformat(saved_summary_dict['dateReleased'].replace("Z", "+00:00"))
    saved_release_str = f"{saved_release_dt:%Y-%m-%d}"  # Reformat this as a simple date string for Discord message

    # Convert string with ISO format date to a datetime
    kev_release_dt = datetime.fromisoformat(kev_data['dateReleased'].replace("Z", "+00:00"))
    kev_release_str = f"{kev_release_dt:%Y-%m-%d}"  # Reformat this as a simple date string for Discord message

    prev_str = (
        f"Date: {saved_release_str}, "
        f"count: {saved_summary_dict['count']},\n "
        f"{vend_name}: {saved_summary_dict['vendCount']}"
    )

    if vend_count == saved_summary_dict['vendCount']:
        # Vendor vuln count has not changed
        vend_count_bold = ""
    else:
        # Vendor vuln count has changed so highlight this in bold
        vend_count_bold = "**"

    latest_str = (
        f"Date: {kev_release_str}, "
        f"count: **{kev_data['count']}**,\n"
        f"{vend_name}: {vend_count_bold}{vend_count}{vend_count_bold}"
    )

    webhook = DiscordWebhook(url=webhook_url)

    embed = DiscordEmbed(title="CISA KEV List updated",
                         color='03b2f8')  # Blue colour for embed

    embed.add_embed_field(name="Previous KEV",
                         value=prev_str, inline=False)
    
    embed.add_embed_field(name="Latest KEV",
                          value=latest_str, inline=False)

    # New vulnerability details — compare against saved CVE IDs for accuracy
    new_vulns = [v for v in vulns_list if v['cveID'] not in saved_cve_ids]
    vulns_str = "\n\n".join(
        v['cveID'] + " *" + v['vendorProject'] + "* - " + v['product']
        for v in new_vulns
    ) or "None identified"

    embed.add_embed_field(name="New vulnerabilities",
                          value=vulns_str, inline=False)

    webhook.add_embed(embed)

    # Submit message to the Discord webhook URL
    response = webhook.execute()

    if response.status_code == 200: # Discord returns 200 for successful webhook posts
        print("Notification to Discord sent successfully.")
    else:
        print(f"Failed to send notification. Status code: {response.status_code}")
        print(response.content)

def notify_vc_to_discord(saved_count, saved_vend_count, vc_data, saved_cve_ids, vend_name, current_vend_count, webhook_url):
    '''Send VulnCheck KEV update notification to Discord'''

    current_count = len(vc_data)

    prev_str = f"count: {saved_count},\n{vend_name}: {saved_vend_count}"

    count_bold      = "**" if saved_count      != current_count      else ""
    vend_count_bold = "**" if saved_vend_count != current_vend_count else ""
    latest_str = (
        f"count: {count_bold}{current_count}{count_bold},\n"
        f"{vend_name}: {vend_count_bold}{current_vend_count}{vend_count_bold}"
    )

    new_vulns = [v for v in vc_data if not (set(v.get("cve") or []) & saved_cve_ids)]
    vulns_str = "\n\n".join(
        ", ".join(v.get("cve") or ["?"]) + " *" + v.get("vendorProject", "") + "* - " + v.get("product", "")
        for v in new_vulns
    ) or "None identified"

    webhook = DiscordWebhook(url=webhook_url)
    embed = DiscordEmbed(title="VulnCheck KEV List updated",
                         color='e74c3c')  # Red colour for embed

    embed.add_embed_field(name="Previous VC KEV",
                          value=prev_str, inline=False)
    embed.add_embed_field(name="Latest VC KEV",
                          value=latest_str, inline=False)
    embed.add_embed_field(name="New vulnerabilities",
                          value=vulns_str, inline=False)

    webhook.add_embed(embed)

    response = webhook.execute()

    if response.status_code == 200:
        print("VulnCheck notification to Discord sent successfully.")
    else:
        print(f"Failed to send VulnCheck notification. Status code: {response.status_code}")
        print(response.content)

def check_kev_updates(kev_header, vulns_list, kev_in_path, notify_discord, webhook_url, vendor_list, vendor_name):
    '''Check if there are any new KEV vulnerabilities since the last saved KEV file'''

    # Convert string with ISO format date to a datetime
    kev_release_dt = datetime.fromisoformat(kev_header['dateReleased'].replace("Z", "+00:00"))
    kev_release_str = f"{kev_release_dt:%Y-%m-%d %H:%M:%S}"  # Reformat this as a more readable string for console output

    # Load existing KEV data from local Excel file
    wb = openpyxl.load_workbook(kev_in_path)
    ws = wb["Summary"]  # Open the summary sheet to compare header info from KEV data

    # Read saved CVE IDs for accurate new-vuln detection
    saved_cve_ids = {row[0] for row in wb["Vulns"].iter_rows(min_row=2, values_only=True) if row[0] is not None}

    # Mapping of spreadsheet header names to KEV header names
    vend_header = "Count of " + vendor_name + " vulns"
    name_dict = {
        "Title": "title",
        "Catalog Version": "catalogVersion",
        "Date Released": "dateReleased",
        "Count of Vulnerabilities": "count",
        vend_header: "vendCount"
    }
    # Get the existing catalog version and date released from the summary sheet of the local workbook
    saved_summary_dict = {}

    for row in ws.iter_rows(min_row=1, values_only=True):
        key = row[0]
        value = row[1]

        if key is not None and key in name_dict:
            saved_summary_dict[name_dict[key]] = value

    # Convert string with ISO format date to a datetime
    saved_release_dt = datetime.fromisoformat(saved_summary_dict['dateReleased'].replace("Z", "+00:00"))
    saved_release_str = f"{saved_release_dt:%Y-%m-%d %H:%M:%S}"  # Reformat this as a more readable string for console output

    vend_count = 0
    for vuln in vulns_list:
        if vuln["vendorProject"] in vendor_list:
            vend_count += 1

    print(f" Saved CISA KEV catalog version: {saved_summary_dict['catalogVersion']}, "
          f"date: {saved_release_str}, "
          f"count: {saved_summary_dict['count']}, "
          f"{vendor_name}: {saved_summary_dict['vendCount']}")

    # Compare catalog versions and dates to determine if KEV list has been updated
    if kev_header["catalogVersion"] != saved_summary_dict['catalogVersion'] or kev_header["dateReleased"] != saved_summary_dict["dateReleased"]:
        if saved_summary_dict['vendCount'] == vend_count:
            # Highlight the overall KEV count in yellow because it has changed
            print(f"Latest CISA KEV catalog version: {Fore.GREEN}{kev_header['catalogVersion']}{Style.RESET_ALL}, "
                  f"date: {Fore.GREEN}{kev_release_str}{Style.RESET_ALL}, "
                  f"count: {Fore.YELLOW}{kev_header['count']}{Style.RESET_ALL}, "
                  f"{vendor_name}: {vend_count}")
        else:
            # Vendor vuln count has changed so highlight this too
            print(f"Latest CISA KEV catalog version: {Fore.GREEN}{kev_header['catalogVersion']}{Style.RESET_ALL}, "
                  f"date: {Fore.GREEN}{kev_release_str}{Style.RESET_ALL}, "
                  f"count: {Fore.YELLOW}{kev_header['count']}{Style.RESET_ALL}, "
                  f"{vendor_name}: {Fore.RED}{vend_count}{Style.RESET_ALL}")

        if notify_discord:
            notify_to_discord(saved_summary_dict, kev_header, vulns_list, vendor_name, vend_count, webhook_url, saved_cve_ids)
        else:
            print("Notification to Discord is not required.")

        return True
    else:
        print(f"Latest CISA KEV catalog version: {kev_header["catalogVersion"]}, "
              f"date: {kev_release_str}, "
              f"count: {kev_header["count"]}, "
              f"{vendor_name}: {vend_count} (no change)")
        
        return False

def save_kev_to_excel(kev_header, vulns_list, vendor_list, vendor_name, kev_in_fn, kev_out_path, vc_meta=None, vc_data=None):
    '''Save KEV data to an Excel spreadsheet'''

    # Mapping of KEV header names to more user-friendly names, for the summary sheet
    name_dict: dict = {
        "title": "Title",
        "catalogVersion": "Catalog Version",
        "dateReleased": "Date Released",
        "count": "Count of Vulnerabilities"
    }

    wb = openpyxl.Workbook()
    
    # Write summary information to the default sheet
    summary_ws = wb.active
    summary_ws.title = "Summary"

    vend_header:str = "Count of " + vendor_name + " vulns"
    for key, value in kev_header.items():
        name = name_dict[key] if key in name_dict else key
        summary_ws.append([name, value])
    summary_ws.append([vend_header, 0])  # Placeholder for vendor vuln count

    # Set column widths for better readability
    summary_ws.column_dimensions['A'].width = 25
    summary_ws.column_dimensions['B'].width = 40

    #Set cells B1 - B5 to bold text
    for row in summary_ws['B1:B5']:
        for cell in row:
            cell.font = styles.Font(bold=True)

    #Set left alignment for the Count cells
    summary_ws['B4'].alignment = styles.Alignment(horizontal='left')
    summary_ws['B5'].alignment = styles.Alignment(horizontal='left')

    # Write vulnerabilities information
    vulns_ws = wb.create_sheet(title="Vulns", index=1)
    # Write header row
    headers = ["CVE ID", "Vendor", "Product", "Vulnerability Name", "Date Added", "Short Description", "Reqd Action", "Due Date", "Ransomware Campaign"]
    vulns_ws.append(headers)

    # Set header row to bold text
    for cell in vulns_ws[1]:
        cell.font = styles.Font(bold=True)

    # Set column widths for better readability
    col_widths = [15, 15, 20, 40, 10, 40, 15, 10, 20]

    for i, width in enumerate(col_widths, start=1):
        vulns_ws.column_dimensions[utils.get_column_letter(i)].width = width

    # Write KEV entries
    for item in vulns_list:
        cve_id: str = item.get("cveID", "")
        vendor: str = item.get("vendorProject", "")
        product: str = item.get("product", "")
        vuln_name: str = item.get("vulnerabilityName", "")
        date_added = item.get("dateAdded", "")
        short_desc:str = item.get("shortDescription", "")
        reqd_action:str = item.get("requiredAction", "")
        due_date = item.get("dueDate", "")
        ransomware_campaign = item.get("knownRansomwareCampaignUse", False)

        row:list = [cve_id, vendor, product, vuln_name, date_added, short_desc, reqd_action, due_date, ransomware_campaign]
        vulns_ws.append(row)

    # Apply autofilter to the vulnerabilities sheet
    vulns_ws.auto_filter.ref = vulns_ws.dimensions

    # Create a separate sheet for vendor vulnerabilities
    vend_vulns_ws = wb.copy_worksheet(vulns_ws)
    vend_vulns_ws.title = vendor_name
    vend_vulns_ws.index = 2

    # Remove non-vendor entries from the vendor sheet
    for row_ctr in range(vend_vulns_ws.max_row, 1, -1):
        vendor_cell = vend_vulns_ws.cell(row=row_ctr, column=2)  # Vendor is in the second column
        if vendor_cell.value not in vendor_list:
            vend_vulns_ws.delete_rows(row_ctr)

    # Apply autofilter to the vendor sheet
    vend_vulns_ws.auto_filter.ref = vend_vulns_ws.dimensions

    # Add the count of vendor entries to the Summary sheet
    vend_count = vend_vulns_ws.max_row - 1
    summary_ws['B5'].value = vend_count

    # Add VulnCheck sheets if VC data was retrieved
    if vc_meta is not None and vc_data is not None:
        write_vc_sheets(wb, vc_meta, vc_data, summary_ws, vendor_list, vendor_name)

    # Save the workbook to the specified file
    today_str = date.today().isoformat()
    out_fn = Path(kev_out_path) / f"CISA_KEV_{today_str}.xlsx"

    wb.save(out_fn)
    print(f"KEV data saved to {out_fn}")

    # Copy the new KEV file to the input file path for future comparisons
    shutil.copy(out_fn, kev_in_fn)

def main():
    # Open config file
    config_path = "config.yaml"
    config_dict = load_config(config_path)

    # Retrieve KEV list from CISA website
    kev_url:str = config_dict['kev']['kev_url']
    kev_data = get_cisa_kev_data(kev_url)

    # Remove the vulnerabilities list from the main dictionary for separate processing
    vulns_list = kev_data.pop("vulnerabilities")
    kev_header = kev_data

    # Path for saved KEV spreadsheet if it exists
    kev_in_fn:str = config_dict['kev']['in_fn']
    kev_in_path = Path(kev_in_fn)
    # Path for output KEV spreadsheet
    kev_out_path:str = config_dict['kev']['out_path']

    # Vendor references from config
    vend_list:list = config_dict['kev']['vendor_list']
    vend_name:str = config_dict['kev']['vendor_name']

    # Optionally fetch VulnCheck KEV data
    use_vc:bool = config_dict['kev'].get('use_vc', False)
    vc_meta, vc_data = None, None

    if use_vc:
        vc_raw = get_vc_kev_data(config_dict['kev']['vc_kev_url'], config_dict['kev']['vc_token'])
        vc_meta = vc_raw.get('_meta', {})
        vc_data = vc_raw.get('data', [])

    if kev_in_path.is_file():   # A saved KEV file exists

        # Check CISA and VulnCheck data independently for changes
        cisa_updated = check_kev_updates(kev_header, vulns_list, kev_in_path, config_dict['kev']['notify'], config_dict['kev']['webhook_url'], vend_list, vend_name)
        vc_updated = check_vc_updates(kev_in_path, vc_data, vend_list, vend_name,
                                      notify_discord=config_dict['kev']['notify'],
                                      webhook_url=config_dict['kev'].get('vc_webhook_url')) if use_vc else False

        if cisa_updated or vc_updated:
            save_kev_to_excel(kev_header, vulns_list, vend_list, vend_name, kev_in_fn, kev_out_path,
                              vc_meta=vc_meta, vc_data=vc_data)
        else:
            print("No changes to KEV data.")

    else:
        # We don't have a previously saved KEV file so we need to save one
        print("Saving KEV data to local Excel spreadsheet...")
        save_kev_to_excel(kev_header, vulns_list, vend_list, vend_name, kev_in_fn, kev_out_path,
                          vc_meta=vc_meta, vc_data=vc_data)

if __name__ == "__main__":
    main()
