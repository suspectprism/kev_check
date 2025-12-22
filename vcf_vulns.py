# Collect VCF vulnerability details from advisories at support.broadcom.com
# and save relevant data into a spreadsheet, maybe a database in future.
#
#
# P Dowley   v0.1      17 Dec 2025

import requests
from bs4 import BeautifulSoup
import re
import pandas as pd
import openpyxl
from config import load_config

def read_sec_advisory_list(in_file):
	'''Read the spreadsheet with a list of security advisories'''
	wb = openpyxl.load_workbook(in_file)
	ws = wb.active # Assume data is in the active sheet

	advisory_URL_list = []

	# Iterate through rows starting from the second row
	for row_num in range(2, ws.max_row + 1):
		cell = ws.cell(row=row_num, column=1) # Advisory ID is in the first column
		try:
			advisory_URL_list.append(cell.hyperlink.target)
		except AttributeError:
			advisory_URL_list.append(None)

	return advisory_URL_list

def process_brcm_advisory(advisory_soup):
	'''Process a Broadcom advisory page and extract vulnerability details'''
	cve_pattern = r'[, ]+' # Regex pattern to split CVEs on commas and spaces (one or more times)

	# Extract details from the page title
	page_title = advisory_soup.find('p', class_='ecx-page-title-default').text.strip()
	title_split = page_title.split(":", 1)

	vmsa = title_split[0].strip() # VMSA identifier

	title_remainder = title_split[1].strip()
	index = title_remainder.find("(")
	if index != -1:
		advisory_title = title_remainder[:index].strip() # Advisory title
	else:
		advisory_title = title_remainder.strip()

	cards = advisory_soup.find_all('div', class_='card-body')

	for card in cards:
		rows = card.find_all('div', class_='row')
		for row in rows:
			cells = row.find_all('div', class_='col-4')
			for cell in cells:
				sol_label = cell.find('label', class_='edit-solution-labels').text.strip()
				sol_text = cell.find('p', class_='edit-solution-text').text.strip()

				if sol_label == "Last Updated":
					upd_date = sol_text
				elif sol_label == "Initial Publication Date":
					pub_date = sol_text
				elif sol_label == "Severity":
					severity = sol_text
				elif sol_label == "CVSS Base Score":
					cvss_score = sol_text
				elif sol_label == "Affected CVE":
					cve_ids = re.split(cve_pattern, sol_text)

	single_cols = advisory_soup.find_all('div', class_='columnLayout single')

	for col in single_cols:
		affected_prods = []

		prod_lines = col.find_all('li')
		for prod_line in prod_lines:
			affected_prods.append(prod_line.text.strip())

	table_df = pd.DataFrame()
	tables = advisory_soup.find_all('tbody')[1:] # Skip the first table which is not relevant

	for table in tables:
		row_list = []
		table_rows = table.find_all('tr')

		for tr in table_rows:
			cell_list = []
			tds = tr.find_all('td')

			for td in tds:
				cell_list.append(td.text.strip())
			
			row_list.append(cell_list)
		
		header = row_list[0] # First row as header
		table_new_df = pd.DataFrame(row_list[1:], columns=header) # Remaining rows as data
		table_df = pd.concat([table_df, table_new_df], ignore_index=True)

	advisory_dict = {
		"VMSA": [vmsa],
		"Synopsis": [advisory_title],
		"Issue date": [pub_date],
		"Update date": [upd_date],
		"Severity": [severity],
		"CVSS Score": [cvss_score]
	}
	#advisory_df = pd.DataFrame(advisory_dict)

	print(advisory_dict)
	print(cve_ids)
	print(affected_prods)
	print(table_df)

	return advisory_dict, cve_ids, affected_prods

def main():
	# Stop chained indexing from being attempted for Pandas pre v3.0
	pd.options.mode.copy_on_write = True

	# Open config file
	config_path = "config.yaml"
	config_dict = load_config(config_path)

	# Get the list of security advisory URLs from the input spreadsheet
	in_file = config_dict['in_fn']
	advisory_list = read_sec_advisory_list(in_file)

	# Fetch a sample VCF advisory page
	response = requests.get(advisory_list[0])
	advisory_content = response.text

	advisory_soup = BeautifulSoup(advisory_content, 'html.parser')

	process_brcm_advisory(advisory_soup)

if __name__ == "__main__":
    main()